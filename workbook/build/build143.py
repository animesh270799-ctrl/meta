"""Post-step of updated 143 (run on the batch.py output): new sheets AGGREGATE STABILITY, AGGREGATE RATIO, AGG-OC CONTRIBUTION for
48 (companion 3) (ext 647, user request), 655 logged as no-eligible-data companion of 211, index sheets rebuilt in workbook order."""
import copy
import re
import sys
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

sys.path.insert(0, sys.path[0])
from wblib import WB
from common import base, notes
import c48

SRC, DST = sys.argv[1], sys.argv[2]
DATE = '2026-10-10'
B = WB(SRC)
wb = B.wb
STYLE_SRC = {}  # new sheet -> (template sheet, template data row, list of template columns per new column)
STD_DEPTH = '"0-15 CM,15-30 CM,30-45 CM,45-60 CM,>60 CM"'


def new_sheet(name, tpl, after, rename, unit, insert=None, extra_lists=None):
    """Copy the header of `tpl`; `insert` = (new header, after header, style-source header) adds one column."""
    t = wb[tpl]
    ws = wb.create_sheet(name, wb.sheetnames.index(after) + 1)
    cols = []  # (header value, template column)
    for c in range(1, t.max_column + 1):
        v = t.cell(1, c).value
        for old, new in rename:
            if isinstance(v, str) and v.startswith(old):
                v = new + v[len(old):]
        cols.append((v, c))
        if insert and t.cell(1, c).value == insert[1]:
            src = [cc for cc in range(1, t.max_column + 1) if t.cell(1, cc).value == insert[2]][0]
            cols.append((insert[0], src))
    for k, (v, c) in enumerate(cols, 1):
        s, d = t.cell(1, c), ws.cell(1, k)
        d.value = v
        d.font = copy.copy(s.font); d.fill = copy.copy(s.fill); d.alignment = copy.copy(s.alignment); d.border = copy.copy(s.border)
        ws.column_dimensions[get_column_letter(k)].width = t.column_dimensions[get_column_letter(c)].width
    ws.row_dimensions[1].height = t.row_dimensions[1].height
    ws.freeze_panes = 'D2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}1'
    ws.cell(2, len(cols)).value = unit
    hdr = {v: get_column_letter(k) for k, (v, _) in enumerate(cols, 1)}
    lists = {'CLIMATE': '"ST,TEMP"', 'SOIL': '"LOAMY,SANDY,CLAYEY"', 'DURATION': '"0-3 Y,4-10 Y,>10 Y"', 'DEPTH': STD_DEPTH}
    lists.update(extra_lists or {})
    for h, f in lists.items():
        dv = DataValidation(type='list', formula1=f, allow_blank=True)
        dv.add(f'{hdr[h]}2:{hdr[h]}1002')
        ws.add_data_validation(dv)
    STYLE_SRC[name] = (tpl, B.template_row(tpl), [c for _, c in cols])
    B._hdr.pop(name, None)
    return ws


new_sheet('AGGREGATE STABILITY', 'GMD 1', 'GMD 1', [('GMD_', 'AS_')],
          'index (aggregate stability, AS = (weight of water-stable aggregates - wp25 - sand) / (weight of dry sample - sand); Castro Filho et al. 2002). Enter as printed; name the formula in UNIT.')
new_sheet('AGGREGATE RATIO', 'GMD 1', 'AGGREGATE STABILITY', [('GMD_', 'AR_')],
          'ratio (aggregate ratio, AR = % water-stable macroaggregates (>0.25 mm) / % water-stable microaggregates (0.25-0.053 mm); Choudhury et al. 2014)')
new_sheet('AGG-OC CONTRIBUTION', 'macro c', 'micro c', [('MACRO c_', 'AOCC_')],
          'g C/kg bulk soil (contribution of one aggregate size class to aggregate-associated OC = class OC x class mass proportion; one row per size class, class in column "Aggregate size class"; % share of total AAOC in Notes)',
          insert=('Aggregate size class', 'DEPTH (as reported in paper)', 'DEPTH (as reported in paper)'),
          extra_lists={'Aggregate size class': '">2 mm,2-0.25 mm,0.25-0.053 mm,<0.053 mm,>0.25 mm,<0.25 mm,1-2 mm,0.5-1 mm,0.25-0.5 mm"'})

_orig_add = B.add_row
NOTE_FILL = copy.copy(wb['TOC'].cell(B.template_row('TOC'), B.hdr('TOC')['Notes/Doubts']).fill)


def add_row(sheet, vals, red=False):
    r = _orig_add(sheet, vals, red)
    if sheet in STYLE_SRC:
        tpl, tr, cols = STYLE_SRC[sheet]
        t, ws = wb[tpl], wb[sheet]
        for k, c in enumerate(cols, 1):
            s, d = t.cell(tr, c), ws.cell(r, k)
            d.border = copy.copy(s.border); d.alignment = copy.copy(s.alignment); d.number_format = s.number_format
            f = copy.copy(s.font)
            if red:
                f = copy.copy(ws.cell(r, k).font) if ws.cell(r, k).font else f
            d.font = f; d.fill = copy.copy(s.fill)
        h = B.hdr(sheet)
        ws.cell(r, h['Notes/Doubts']).fill = copy.copy(NOTE_FILL)
        if red:
            for k in range(1, len(cols) + 1):
                cell = ws.cell(r, k); fnt = copy.copy(cell.font); fnt.color = 'FFFF0000'; fnt.name = 'Arial'; cell.font = fnt
    return r


B.add_row = add_row

# ---------------------------------------------------------------- 48 (companion 3): Fig. 2c / 2d and Fig. 3
st, SITE, C3, RED, HEAD = c48.ST, c48.SITE, c48.C3, c48.RED, c48.HEAD
UPD = 'ADDED updated 143 (user request 2026-10-10, new sheet). '
MA = c48.__dict__.get('ma') or ('Yoder wet sieving (2, 0.25, 0.053 mm sieves; 50 g <8 mm air-dried soil, capillary-wetted 10 min, 15 min at 35 cycles/min), sand-corrected.')
LY2 = [('0-10', '0-15 CM'), ('10-20', '15-30 CM')]
cnt_rows = []


def add(sheet, vals, body, depth, unit, src, meth, extra=None):
    obs, nt = notes(st, UPD + HEAD + ' ' + body)
    v = base(st, SITE, {'year of data collection/experiment': '2014', 'YEAR OF DATA (duration)': '2014 (5 yr)'})
    v.update({'DEPTH': depth[1], 'DEPTH (as reported in paper)': depth[0] + ' cm'})
    v.update(extra or {})
    v.update(vals)
    v.update({'UNIT': unit, 'Crop/season of sampling': RED, 'Data source': src, 'Method used (from paper)': meth, 'Obs': obs, 'Rep': 3, 'Notes/Doubts': nt})
    cnt_rows.append((sheet, B.add_row(sheet, v, red=True)))


# Fig. 2c / 2d digitised (raster, calibrated axis ticks, bar-top outline +1 px); order TA, pCA1, fCA, pCA2
AS = {'0-10': (1.327, 1.498, 1.684, 1.489), '10-20': (1.317, 1.532, 1.669, 1.489)}
AR = {'0-10': (3.65, 6.16, 7.91, 5.17), '10-20': (3.25, 4.68, 7.91, 4.62)}
T2 = {'0-10': [(66.4, 18.5), (74.9, 12.8), (83.7, 10.7), (74.0, 14.6)], '10-20': [(65.4, 20.2), (76.2, 16.8), (83.3, 13.0), (74.0, 16.5)]}
TXT = {'0-10': ('25.6', '119'), '10-20': ('27.5', '144')}
for ly, cls in LY2:
    a = AS[ly]; r = AR[ly]
    add('AGGREGATE STABILITY', {f'AS_{c}': x for c, x in zip(C3, a[:3])},
        f'Figure 2c aggregate stability {ly} cm, DIGITISED (calibrated axis 0-2.0; error bars not read). pCA2 {a[3]}. Text check: fCA {TXT[ly][0]} % higher than TA (digitised {round((a[2] / a[0] - 1) * 100, 1)} %). '
        'Eq. 3 prints AS x 100 but the plotted values are 1.3-1.7 (index scale as printed - FLAG).',
        (ly, cls), 'index (AS, Castro Filho et al. 2002; as plotted)', 'Figure 2c (digitised)', MA + ' AS = (aggregates - wp25 - sand) / (dry sample - sand) (Eq. 3).')
    ratio = '; '.join(f'{c} {m} / {mi} = {round(m / mi, 2)}' for c, (m, mi) in zip(C3 + ['pCA2'], T2[ly]))
    add('AGGREGATE RATIO', {f'AR_{c}': x for c, x in zip(C3, r[:3])},
        f'Figure 2d aggregate ratio {ly} cm, DIGITISED (calibrated axis 0-12; error bars not read). pCA2 {r[3]}. Text check: fCA {TXT[ly][1]} % higher than TA (digitised {round((r[2] / r[0] - 1) * 100)} %). '
        f'Ratio of the Table 2 MEANS (macro / 0.25-0.053 mm) differs where replicate ratios vary (mean of ratios plotted): {ratio}.',
        (ly, cls), 'ratio (% macroaggregates / % microaggregates, as plotted)', 'Figure 2d (digitised)', MA + ' AR = % water-stable macroaggregates / % water-stable microaggregates (Eq. 4).')

# Fig. 3 stacked bars, digitised (segment boundaries; g/100 g), order TA, pCA1, fCA, pCA2: (MacOC, MesOC, MicOC, SCOC)
F3 = {'0-10': [(0.078, 0.312, 0.126, 0.039), (0.256, 0.347, 0.084, 0.052), (0.302, 0.299, 0.091, 0.049), (0.333, 0.274, 0.062, 0.032)],
      '10-20': [(0.070, 0.290, 0.121, 0.042), (0.209, 0.221, 0.083, 0.031), (0.286, 0.147, 0.076, 0.032), (0.241, 0.181, 0.053, 0.017)]}
TAOC = {'0-10': (0.55, 0.74, 0.74, 0.70), '10-20': (0.53, 0.54, 0.54, 0.49)}
CLS = [('MacOC', '8-2 mm', '>2 mm'), ('MesOC', '2-0.25 mm', '2-0.25 mm'), ('MicOC', '0.25-0.053 mm', '0.25-0.053 mm'), ('SCOC', '<0.053 mm (silt + clay)', '<0.053 mm')]
for ly, cls in LY2:
    f = F3[ly]
    tot = [sum(x) for x in f]
    for j, (lab, rep, sc) in enumerate(CLS):
        share = '; '.join(f'{c} {round(100 * x[j] / t, 1)} %' for c, x, t in zip(C3 + ['pCA2'], f, tot))
        add('AGG-OC CONTRIBUTION', {f'AOCC_{c}': f'=ROUND({x[j]}*10,2)' for c, x in zip(C3, f[:3])},
            f'Figure 3{"a" if ly == "0-10" else "b"} contribution of {lab} ({rep}) to aggregate-associated OC, {ly} cm, DIGITISED stacked segment (g/100 g x 10). pCA2 {f[3][j]} g/100 g. '
            f'Share of total AAOC: {share}. Digitised stack totals {" / ".join(str(round(t, 3)) for t in tot)} vs Table 4 tAOC {" / ".join(map(str, TAOC[ly]))} g/100 g (check). '
            'Note: Fig. 3 segments do not equal Table 2 class proportion x Table 4 class OC (paper\'s calculation not given) - FLAG.',
            (ly, cls), 'g C/kg bulk soil (printed g/100 g x 10; DIGITISED)', 'Figure 3 (digitised)',
            'Aggregate-associated OC (Walkley-Black) of each wet-sieved size class weighted by its mass proportion; contribution to total AAOC.',
            extra={'Aggregate size class': sc})

# existing 48 (companion 3) notes / Study_Info / PARAMETERS
def fix_notes(sheet, serial, pattern, repl):
    ws = wb[sheet]; h = B.hdr(sheet); n = 0
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, h['SERIAL NO']).value) != serial:
            continue
        c = ws.cell(r, h['Notes/Doubts'])
        if c.value and re.search(pattern, c.value):
            c.value = re.sub(pattern, repl, c.value); n += 1
    return n


n = fix_notes('MWD 1', '48 (companion 3)', r' Aggregate stability \(Fig\. 2c[^)]*\) [^;]*; aggregate ratio \(Fig\. 2d\) [^-]*- no sheet\.',
              ' Aggregate stability (Fig. 2c) and aggregate ratio (Fig. 2d): on the AGGREGATE STABILITY / AGGREGATE RATIO sheets (updated 143).')
print('MWD notes updated:', n)
si = wb['Study_Info']; h = B.hdr('Study_Info')
for r in range(2, si.max_row + 1):
    if str(si.cell(r, h['SERIAL NO']).value) == '48 (companion 3)':
        nv = si.cell(r, h['Notes/Doubts']).value
        nv = nv.replace('Not entered (no sheet): aggregate stability, aggregate ratio, % contribution of aggregate classes to AOC (Fig. 3).',
                        'UPDATE updated 143 (user request): aggregate stability, aggregate ratio (Fig. 2c-d) and aggregate-class contributions to AAOC (Fig. 3) entered on the new sheets '
                        'AGGREGATE STABILITY, AGGREGATE RATIO and AGG-OC CONTRIBUTION.')
        si.cell(r, h['Notes/Doubts']).value = nv
        si.cell(r, h['Parameters extracted']).value += '; [updated 143] aggregate stability, aggregate ratio (Fig. 2c-d), aggregate-class contribution to AAOC (Fig. 3) - digitised, 0-10 / 10-20 cm'

# ---------------------------------------------------------------- 655: no eligible data (companion trial of 211)
B.add_row('EXCLUDED_rows', {'Sheet': 'ALL', 'SERIAL NO': 'ext\\655.pdf = trial of 211 (not entered)', 'Authors': 'Shahzad M., Farooq M. & Hussain M.', 'Year': 2016,
                            'Reason for exclusion': ('EXCLUDED - NO ELIGIBLE DATA (rule 136): Shahzad M., Farooq M. & Hussain M. (2016) Weed spectrum in different wheat-based cropping systems under conservation and '
                                                     'conventional tillage practices in Punjab, Pakistan. Soil & Tillage Research 163:71-79, doi 10.1016/j.still.2016.05.012 - same BZU Multan trial as study 211 / 211 (companion) '
                                                     '(2012-14; ZT / CT / DT / beds x 5 wheat-based systems). Only weed diversity and weed densities by species (Fig. 2, Tables 3-6) - no sheet. '
                                                     'Rice-wheat cells 2012-13 / 2013-14 total weeds per m2: ZT 87.00 / 86.67, CT 26.00 / 20.67, DT 15.33 / 12.67, BS 60/30 11.67 / 9.00, BS 90/45 11.00 / 7.33.'),
                            'Full row (header = value)': 'No rows.'})

# ---------------------------------------------------------------- index sheets (rebuilt in workbook order)
names = wb.sheetnames
data_sheets = names[names.index('LAT_LONG'):names.index('EXCLUDED_rows')]
cnt = {}
for nme in data_sheets:
    ws = wb[nme]; bc = B.hdr(nme)['SERIAL NO']; c = {}
    for r in range(2, ws.max_row + 1):
        val = ws.cell(r, bc).value
        if val is None:
            continue
        k = str(val).strip(); c[k] = c.get(k, 0) + 1
    cnt[nme] = c
sp = wb['STUDIES_BY_PARAMETER']
units = {sp.cell(r, 2).value: sp.cell(r, 3).value for r in range(4, sp.max_row + 1)}
units.update({'AGGREGATE STABILITY': 'index', 'AGGREGATE RATIO': 'ratio', 'AGG-OC CONTRIBUTION': 'g C/kg bulk soil'})
tpl = 5
for i, nme in enumerate(data_sheets):
    r = 4 + i
    for c in range(1, 7):
        s_, d_ = sp.cell(tpl, c), sp.cell(r, c)
        d_.font = copy.copy(s_.font); d_.alignment = copy.copy(s_.alignment); d_.border = copy.copy(s_.border); d_.fill = copy.copy(s_.fill)
    c = cnt[nme]
    vals = [i + 1, nme, units.get(nme), sum(c.values()), len(c), '; '.join(f'{k}: {v}' for k, v in c.items()) or 'NO DATA']
    for j, v in enumerate(vals, 1):
        sp.cell(r, j).value = v
ps = wb['PARAMETERS_BY_STUDY']
for r in range(4, ps.max_row + 1):
    s = str(ps.cell(r, 2).value)
    if s == '48 (companion 3)':
        parts = [(nme, cnt[nme][s]) for nme in data_sheets if s in cnt[nme]]
        ps.cell(r, 7).value = sum(v for _, v in parts); ps.cell(r, 8).value = len(parts)
        ps.cell(r, 9).value = '; '.join(f'{k}: {v}' for k, v in parts)
        for rr2 in range(2, si.max_row + 1):
            if str(si.cell(rr2, h['SERIAL NO']).value) == s:
                ps.cell(r, 10).value = si.cell(rr2, h['Parameters extracted']).value

# restore the orange unit note style in row 2 of the new sheets
for nm, (tpl_, _, cols) in STYLE_SRC.items():
    ws = wb[nm]; t = wb[tpl_]
    s2 = t.cell(2, t.max_column); d2 = ws.cell(2, len(cols))
    d2.font = copy.copy(s2.font); d2.fill = copy.copy(s2.fill); d2.alignment = copy.copy(s2.alignment); d2.border = copy.copy(s2.border)
wb.calculation.fullCalcOnLoad = True
wb.save(DST)
from collections import Counter
print('rows added per sheet:', dict(Counter(s for s, _ in B.log)))
