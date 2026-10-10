import openpyxl, sys
a=openpyxl.load_workbook(sys.argv[1],read_only=True); b=openpyxl.load_workbook(sys.argv[2],read_only=True)
assert a.sheetnames==b.sheetnames, 'sheet list changed'
for n in a.sheetnames:
    ra=list(a[n].iter_rows(values_only=True)); rb=list(b[n].iter_rows(values_only=True))
    if n=='Bibliography':
        sa=set(r[1:] for r in ra if r and r[0]!=None and isinstance(r[0],int)); sb=set(r[1:] for r in rb if r and isinstance(r[0],int))
        print(n,'missing old entries:',len(sa-sb),'new entries:',len(sb-sa)); continue
    ch=0; first=[]
    for i,(x,y) in enumerate(zip(ra,rb),1):
        x=tuple(x)+(None,)*(len(y)-len(x)) if len(x)<len(y) else tuple(x)
        y=tuple(y)+(None,)*(len(x)-len(y)) if len(y)<len(x) else tuple(y)
        for j,(p,q) in enumerate(zip(x,y),1):
            if p!=q and not (p is None and n not in ("STUDIES_BY_PARAMETER",)):
                ch+=1
                if len(first)<3: first.append((i,j,str(p)[:40],str(q)[:40]))
    extra=len(rb)-len(ra)
    if ch or extra: print(n,'changed cells in old rows:',ch,'new rows:',extra, first if ch else '')
