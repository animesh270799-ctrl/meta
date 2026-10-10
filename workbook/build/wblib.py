"""Helpers to append rows to META_ANALYSIS_MASTER in the workbook's own conventions."""
import copy
import openpyxl
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter

NOTES_FILL = PatternFill(fill_type='solid', fgColor='FFFFF2CC', bgColor='FFFFF2CC')
RED = 'FFFF0000'

COMMON_SITE_KEYS = ['No.', 'SERIAL NO', 'Authors', 'Year', 'Journal', 'Country', 'Site/Location',
                    'latitude', 'longitude', 'CLIMATE', 'year of data collection/experiment',
                    'DURATION', 'SOIL']
COVARIATES = ['CLAY', 'MIN TEMP', 'MAX TEMP', 'RAIN FALL', 'LATT', 'YEAR OF DATA (duration)', 'AVG T',
              'April max temp', 'ph (initial)', 'soc (initial)', 'Bdi', 'sand', 'silt']


class WB:
    def __init__(self, path):
        self.wb = openpyxl.load_workbook(path)
        self._hdr = {}
        self.log = []  # (sheet, row)

    def hdr(self, sheet):
        if sheet not in self._hdr:
            ws = self.wb[sheet]
            m = {}
            for c in ws[1]:
                if c.value is not None:
                    k = str(c.value).strip()
                    if k not in m:
                        m[k] = c.column
            self._hdr[sheet] = m
        return self._hdr[sheet]

    def last_row(self, sheet):
        ws = self.wb[sheet]
        col = self.hdr(sheet).get('SERIAL NO', 2)
        r = ws.max_row
        while r > 1 and ws.cell(r, col).value in (None, ''):
            r -= 1
        return r

    def template_row(self, sheet):
        """Last data row that is not red (style source)."""
        ws = self.wb[sheet]
        r = self.last_row(sheet)
        rr = r
        while rr > 1:
            f = ws.cell(rr, 1).font
            if not (f.color is not None and f.color.rgb == RED):
                return rr
            rr -= 1
        return r

    def col(self, sheet, name):
        h = self.hdr(sheet)
        if name not in h:
            raise KeyError(f'{sheet}: no header {name!r}')
        return h[name]

    def letter(self, sheet, name):
        return get_column_letter(self.col(sheet, name))

    def add_row(self, sheet, vals, red=False):
        ws = self.wb[sheet]
        h = self.hdr(sheet)
        tpl = self.template_row(sheet)
        r = self.last_row(sheet) + 1
        maxc = max(h.values())
        for c in range(1, maxc + 1):
            src = ws.cell(tpl, c)
            dst = ws.cell(r, c)
            if src.has_style:
                dst.font = copy.copy(src.font)
                dst.border = copy.copy(src.border)
                dst.alignment = copy.copy(src.alignment)
                dst.number_format = src.number_format
                dst.fill = copy.copy(src.fill)
        for k, v in vals.items():
            if v is None or v == '':
                continue
            ws.cell(r, self.col(sheet, k)).value = v
        # SD formula
        if 'Obs' in h and 'Rep' in h and 'SD' in h and 'SD' not in vals:
            o = get_column_letter(h['Obs']); p = get_column_letter(h['Rep'])
            ws.cell(r, h['SD']).value = (f'=IF(AND(N({o}{r})>0,N({p}{r})>0),'
                                         f'SQRT(1/((({p}{r}*{p}{r})/({p}{r}+{p}{r}))/{o}{r})),"")')
        if 'Notes/Doubts' in h:
            fill = NOTES_FILL if sheet != 'Study_Info' else PatternFill(fill_type='solid', fgColor='FFFFF3CD', bgColor='FFFFF3CD')
            ws.cell(r, h['Notes/Doubts']).fill = copy.copy(fill)
        for c in range(1, maxc + 1):
            cell = ws.cell(r, c)
            f = copy.copy(cell.font) if cell.has_style else Font(name='Arial', size=10)
            if f.name != 'Arial':
                f = Font(name='Arial', size=10, bold=f.bold, italic=f.italic, color=f.color)
            if red:
                f = Font(name=f.name, size=f.size, bold=f.bold, italic=f.italic, color=RED)
            elif f.color is not None and f.color.rgb == RED:
                f = Font(name=f.name, size=f.size, bold=f.bold, italic=f.italic)
            cell.font = f
        self.log.append((sheet, r))
        return r
