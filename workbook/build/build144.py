"""Post-step of updated 144: RULES row 142 (study-specific author coding decisions 2026-10-10)."""
import copy
import sys
import openpyxl
SRC, DST = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC)
ws = wb['RULES']
r = ws.max_row
while ws.cell(r, 1).value is None:
    r -= 1
n = r + 1
for c in range(1, 7):
    s, d = ws.cell(r, c), ws.cell(n, c)
    d.font = copy.copy(s.font); d.alignment = copy.copy(s.alignment); d.fill = copy.copy(s.fill); d.border = copy.copy(s.border)
vals = [142, 'Treatment codes',
        ('STUDY-SPECIFIC CODING (author 2026-10-10): 326 (ext 651) wheat T2 cultivator + rotavator once each -> MT (paper groups it with minimum tillage; not CT under the rotary rule, not pMT); '
         '328 (ext 653) conventional-wheat MAIN EFFECTS entered as CT although pooled over the excluded zero-till DSR + conventional wheat cell (exception to rule 65 for this study); '
         '329 (ext 654) T2 reduced-tilled unpuddled transplanted rice + zero-till wheat -> ZT (row c) (not MT under rule 127, not pZT under rule 138); '
         '327 (ext 652) 125 % N residue rows (CA rows b / c, flagged) kept.'),
        'Applied in updated 144 (326 / 328 / 329 rebuilt; values unchanged, codes and rows as stated). Rules 13, 65, 124, 127, 138 are unchanged elsewhere unless the author asks.',
        'Author 2026-10-10', 'IN FORCE (author)']
for c, v in enumerate(vals, 1):
    ws.cell(n, c).value = v
wb.calculation.fullCalcOnLoad = True
wb.save(DST)
print('RULES row', n)
