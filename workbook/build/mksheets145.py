"""Create the ACTIVE C POOL and PASSIVE C POOL sheets (author request 2026-10-11) after NLC(Cfrac4); row 2 pre-styled from the template data row."""
import copy
import sys
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
SRC, DST = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC)
RED = 'FFFF0000'


def tpl_row(ws):
    h = {c.value: c.column for c in ws[1]}
    col = h['SERIAL NO']
    r = ws.max_row
    while r > 1 and ws.cell(r, col).value in (None, ''):
        r -= 1
    rr = r
    while rr > 1:
        f = ws.cell(rr, 1).font
        if not (f.color is not None and f.color.rgb == RED):
            return rr
        rr -= 1
    return r


def new_sheet(name, tpl, after, old, new, unit):
    t = wb[tpl]
    ws = wb.create_sheet(name, wb.sheetnames.index(after) + 1)
    n = t.max_column
    tr = tpl_row(t)
    for c in range(1, n + 1):
        s, d = t.cell(1, c), ws.cell(1, c)
        v = s.value
        if isinstance(v, str) and v.startswith(old):
            v = new + v[len(old):]
        d.value = v
        d.font = copy.copy(s.font); d.fill = copy.copy(s.fill); d.alignment = copy.copy(s.alignment); d.border = copy.copy(s.border)
        ws.column_dimensions[get_column_letter(c)].width = t.column_dimensions[get_column_letter(c)].width
        if c < n:
            s2, d2 = t.cell(tr, c), ws.cell(2, c)
            d2.font = copy.copy(s2.font); d2.fill = copy.copy(s2.fill); d2.alignment = copy.copy(s2.alignment); d2.border = copy.copy(s2.border); d2.number_format = s2.number_format
    s2, d2 = t.cell(2, n), ws.cell(2, n)
    d2.value = unit
    d2.font = copy.copy(s2.font); d2.fill = copy.copy(s2.fill); d2.alignment = copy.copy(s2.alignment); d2.border = copy.copy(s2.border)
    ws.row_dimensions[1].height = t.row_dimensions[1].height
    ws.freeze_panes = 'D2'
    ws.auto_filter.ref = f'A1:{get_column_letter(n)}1'
    for dv in t.data_validations.dataValidation:
        nd = DataValidation(type=dv.type, formula1=dv.formula1, allow_blank=True)
        nd.sqref = dv.sqref
        ws.add_data_validation(nd)
    print(name, 'created from', tpl, 'template row', tr)


new_sheet('ACTIVE C POOL', 'NLC(Cfrac4)', 'NLC(Cfrac4)', 'NLC_', 'ACT_',
          'g/kg (ACTIVE / labile C pool = very labile + labile C, Chan et al. Fract. 1 + Fract. 2 (12 N and 18 N H2SO4-oxidisable C); state the definition in UNIT)')
new_sheet('PASSIVE C POOL', 'NLC(Cfrac4)', 'ACTIVE C POOL', 'NLC_', 'PAS_',
          'g/kg (PASSIVE / non-labile C pool = less labile + recalcitrant C, Chan et al. Fract. 3 + Fract. 4; state the definition in UNIT)')
wb.save(DST)
