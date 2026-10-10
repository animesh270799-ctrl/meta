import openpyxl, collections
wb=openpyxl.load_workbook('master.xlsx',read_only=True)
names=wb.sheetnames
start=names.index('LAT_LONG'); end=names.index('EXCLUDED_rows')
data=names[start:end]
cnt={}
for n in data:
    ws=wb[n]; c=collections.OrderedDict()
    hdr=[x for x in next(ws.iter_rows(min_row=1,max_row=1,values_only=True))]
    bi=hdr.index('SERIAL NO')
    for row in ws.iter_rows(min_row=2,values_only=True):
        if row[bi] is None: continue
        k=str(row[bi]).strip()
        c[k]=c.get(k,0)+1
    cnt[n]=c
ps=wb['PARAMETERS_BY_STUDY']
mism=0
for row in ps.iter_rows(min_row=4,values_only=True):
    if row[1] is None: continue
    s=str(row[1]).strip()
    parts=[f"{n}: {cnt[n][s]}" for n in data if s in cnt[n]]
    mine='; '.join(parts) if parts else '-'
    if mine!=row[8]:
        mism+=1
        if mism<6: print('MISMATCH',s,'\n mine:',mine[:300],'\n file:',str(row[8])[:300], row[6], sum(cnt[n][s] for n in data if s in cnt[n]))
print('mismatches',mism)
sp=wb['STUDIES_BY_PARAMETER']
m2=0
for row in sp.iter_rows(min_row=4,values_only=True):
    if row[1] is None: continue
    n=row[1]
    if n not in cnt: print('not data sheet',n); continue
    mine='; '.join(f"{k}: {v}" for k,v in cnt[n].items()) or 'NO DATA'
    if mine!=row[5]:
        m2+=1
        if m2<4: print('SP MISMATCH',n,mine[:200],'|',str(row[5])[:200])
print('sp mismatches',m2, [r[1] for r in sp.iter_rows(min_row=4,max_row=8,values_only=True)])
print(data[:5], len(data))
